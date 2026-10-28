"""Generate a deliberately varied demonstration DOCX, not a research result."""
import io
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from lxml import etree as E
from engine import formula, q

def create_sample():
    d=Document()
    d.add_heading('企业创新与金融发展',0)
    d.add_paragraph('演示文档 · 用于检查标题、段落、复杂表格和原生公式')
    d.add_paragraph('摘要：本文件展示 DocStyle AI 的格式编辑功能。表格数值均为演示数据，不代表真实研究结果。')
    d.add_heading('一、研究设计',1)
    p=d.add_paragraph('这段正文包含需要保留的强调：')
    p.add_run('原文文字不应改变').bold=True
    p.add_run('。你可以改变字体、首行缩进、行距及段前段后间距，预览会同步更新。')
    d.add_heading('1.1 变量与样本',2)
    d.add_paragraph('选择段落后，可修正其结构类型，并为当前段落单独设置格式。')
    d.add_heading('1.1.1 指标定义',3)
    d.add_paragraph('净资产收益率、公司规模和杠杆率作为示例指标。英文与数字使用独立字体设置。')
    d.add_heading('1.1.1.1 补充说明',4)
    d.add_paragraph('第 4–9 级标题同样可以独立定义字体、字号、间距和分页行为。')
    d.add_paragraph('表 1 演示指标与分组统计',style='Caption')
    t=d.add_table(rows=5,cols=4); t.style='Table Grid'
    t.cell(0,0).merge(t.cell(1,0)).text='指标'
    t.cell(0,1).merge(t.cell(0,2)).text='演示统计'
    t.cell(0,3).merge(t.cell(1,3)).text='说明'
    t.cell(1,1).text='均值'; t.cell(1,2).text='标准差'
    for row,values in zip(t.rows[2:],[['Size','22.14','1.25','对数值'],['ROA','0.052','0.030','示例'],['Lev','0.46','0.18','示例']]):
        for cell,value in zip(row.cells,values): cell.text=value
    d.add_paragraph('注：合并表头与数据应在修改边框、列宽和单元格格式后保持不变。')
    d.add_heading('二、模型设定',1)
    d.add_paragraph('以下为可在 Word 中继续编辑的原生公式：')
    p=d.add_paragraph(); omath,_=formula(r'y_{it}=\alpha+\beta x_{it}+\gamma Z_{it}+\mu_i+\lambda_t+\varepsilon_{it}')
    mp=E.SubElement(p._p,q('m:oMathPara')); mp.append(omath)
    d.add_paragraph('下面的段落包含行内公式：')
    p=d.add_paragraph('对于 '); math,_=formula(r'\frac{a}{b}'); p._p.append(math); p.add_run('，原有文字与公式对象均保留。')
    d.add_heading('参考文献',1)
    d.add_paragraph('[1] 作者. 文献标题（示例）[J]. 示例期刊, 2026.')
    section=d.sections[0]; section.header.paragraphs[0].text='DocStyle AI · 演示样本'
    section.footer.paragraphs[0].text='演示文档，仅用于功能验证'
    out=io.BytesIO(); d.save(out); return out.getvalue()

if __name__=='__main__':
    Path(__file__).with_name('演示文档.docx').write_bytes(create_sample())
