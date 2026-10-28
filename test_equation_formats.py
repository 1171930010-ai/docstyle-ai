"""Formula formatting: self-built OOXML fixtures; no real MathType claims."""
import copy, io, unittest
from lxml import etree as E
from docx import Document
from engine import NS, child, default_profile, formula, inspect_doc, package, parse, q, transform, val
from sample import create_sample


def profile_without_text_changes():
    p=default_profile()
    for s in p['styles'].values(): s['enabled']=False
    return p


def standalone_sources():
    d=Document();d.add_paragraph('正文不改变。');d.add_paragraph(r'$$\frac{a_1}{b}=2$$');d.add_paragraph(r'普通文本 $x$ 保持原样')
    out=io.BytesIO();d.save(out);return out.getvalue()


def ole_fixture():
    d=Document(); p=d.add_paragraph(); run=p.add_run()._r
    obj=E.SubElement(run,q('object'))
    E.SubElement(obj,'{urn:schemas-microsoft-com:office:office}OLEObject',ProgID='Equation.DSMT4')
    out=io.BytesIO();d.save(out);return out.getvalue()

class EquationFormatsTests(unittest.TestCase):
    def test_latex_delimiters_and_unsafe_commands(self):
        for src in [r'$x_1$',r'$$x_1$$',r'\(x_1\)',r'\[x_1\]']:
            converted=''.join(formula(src)[0].itertext())
            self.assertEqual(converted,'x1')
        with self.assertRaises(ValueError): formula(r'$\href{https://example.com}{x}$')

    def test_global_size_local_override_preserves_formula_content(self):
        source=create_sample();p=profile_without_text_changes()
        p['equation'].update(enabled=True,size=18,scope='all',_fieldStates={'size':'set','align':'preserve','before':'preserve','after':'preserve'})
        result,report=transform(source,{'profile':p,'equations':{'e0':{'format':{'size':14,'_fieldStates':{'size':'set'}}}}})
        before=parse(package(source)[0]['word/document.xml']);after=parse(package(result)[0]['word/document.xml'])
        b=before.findall('.//'+q('m:oMath'));a=after.findall('.//'+q('m:oMath'))
        self.assertEqual([[x.text for x in n.findall('.//'+q('m:t'))] for n in b],[[x.text for x in n.findall('.//'+q('m:t'))] for n in a])
        for idx,expected in [(0,'28'),(1,'36')]:
            runs=a[idx].findall('.//'+q('m:r'));self.assertTrue(runs)
            self.assertTrue(all(val(child(child(r,'rPr'),'sz'))==expected for r in runs))
        self.assertGreater(report['counts']['equationsFormatted'],0)

    def test_unspecified_fields_do_not_mutate_original_formula(self):
        source=create_sample();p=profile_without_text_changes();p['equation'].update(enabled=True,_fieldStates={'align':'preserve','before':'preserve','after':'preserve','size':'unspecified','font':'unspecified'})
        result,_=transform(source,{'profile':p});before=parse(package(source)[0]['word/document.xml']);after=parse(package(result)[0]['word/document.xml'])
        self.assertEqual([E.tostring(x,method='c14n') for x in before.findall('.//'+q('m:oMath'))],[E.tostring(x,method='c14n') for x in after.findall('.//'+q('m:oMath'))])

    def test_latex_candidate_requires_explicit_confirmation(self):
        source=standalone_sources();model=inspect_doc(source)
        candidates=[x for x in model['equations'] if x.get('kind')=='latex_text']
        self.assertEqual(len(candidates),1);self.assertEqual(candidates[0]['latex'],r'\frac{a_1}{b}=2')
        p=profile_without_text_changes();unchanged,_=transform(source,{'profile':p})
        self.assertIn(r'$$\frac{a_1}{b}=2$$', ''.join(Document(io.BytesIO(unchanged)).paragraphs[1].text))
        result,report=transform(source,{'profile':p,'equations':{candidates[0]['id']:{'convertLatex':True}}})
        after=inspect_doc(result)
        self.assertTrue(any(x.get('kind')=='native' for x in after['equations']))
        self.assertIn('普通文本 $x$ 保持原样',[x['text'] for x in after['paragraphs']]);self.assertEqual(report['counts']['latexConverted'],1)

    def test_mathtype_object_detected_and_preserved_when_layout_changes(self):
        source=ole_fixture();model=inspect_doc(source);obj=next(x for x in model['equations'] if x.get('kind')=='mathtype_ole')
        self.assertEqual(obj['id'],'o0');p=profile_without_text_changes();p['equation'].update(enabled=True,align='right',before=7,after=8)
        result,report=transform(source,{'profile':p});before=parse(package(source)[0]['word/document.xml']);after=parse(package(result)[0]['word/document.xml'])
        self.assertEqual(E.tostring(before.find('.//'+q('object')),method='c14n'),E.tostring(after.find('.//'+q('object')),method='c14n'))
        self.assertEqual(val(child(child(after.find('.//'+q('p')),'pPr'),'jc')),'right')
        self.assertTrue(any('MathType' in x for x in report['warnings']))

if __name__=='__main__': unittest.main()

class ExtendedEquationTests(unittest.TestCase):
    def test_font_color_bold_and_italic_preserve_unrelated_math_styles(self):
        source=create_sample(); p=profile_without_text_changes()
        before=parse(package(source)[0]['word/document.xml']);math=before.findall('.//'+q('m:oMath'))[0]
        runs=math.findall('.//'+q('m:r')); original_text=[r.find(q('m:t')).text for r in runs]
        result,report=transform(source,{'profile':p,'equations':{'e0':{'format':{'font':'Cambria Math','size':15.5,'bold':True,'color':'345678'}}}})
        root=parse(package(result)[0]['word/document.xml']);changed=root.findall('.//'+q('m:oMath'))[0]
        self.assertEqual(original_text,[r.find(q('m:t')).text for r in changed.findall('.//'+q('m:r'))])
        for r in changed.findall('.//'+q('m:r')):
            self.assertEqual(val(child(child(r,'rPr'),'sz')),'31')
            self.assertEqual(val(child(child(r,'rPr'),'rFonts'),'ascii'),'Cambria Math')
            self.assertEqual(val(child(child(r,'rPr'),'color')),'345678')
            self.assertEqual(val(child(child(r,'rPr'),'b')),'1')
        self.assertEqual(report['counts']['equationsFormatted'],1)
        with self.assertRaisesRegex(ValueError,'0.5'):
            transform(source,{'profile':p,'equations':{'e0':{'format':{'size':15.3}}}})

    def test_global_scope_excludes_inline_and_global_applies_to_new_formula(self):
        source=create_sample();p=profile_without_text_changes()
        p['equation'].update(enabled=True,scope='display',size=17,_fieldStates={'size':'set','align':'preserve','before':'preserve','after':'preserve'})
        result,report=transform(source,{'profile':p,'insertions':[{'id':'n1','after':None,'latex':r'a^2+b^2=c^2'}]})
        math=parse(package(result)[0]['word/document.xml']).findall('.//'+q('m:oMath'))
        self.assertEqual(val(child(child(math[0].find('.//'+q('m:r')),'rPr'),'sz')),'34')
        self.assertIsNone(child(child(math[1].find('.//'+q('m:r')),'rPr'),'sz'))
        self.assertEqual(val(child(child(math[2].find('.//'+q('m:r')),'rPr'),'sz')),'34')
        self.assertEqual(report['counts']['equationsFormatted'],2)

    def test_latex_conversion_with_normal_profile_and_format(self):
        source=standalone_sources();p=default_profile();p['equation'].update(enabled=True,size=15,_fieldStates={'size':'set','align':'preserve','before':'preserve','after':'preserve'})
        result,report=transform(source,{'profile':p,'equations':{'l0':{'convertLatex':True,'layout':{'after':9}}}})
        new=inspect_doc(result);converted=next(x for x in new['equations'] if x['kind']=='native')
        self.assertTrue(all(f.get('size')==15 for f in converted['runFormats']))
        self.assertEqual(next(x for x in new['paragraphs'] if x['id']==converted['paragraph'])['format']['after'],9)
        self.assertEqual(report['latexConversions'][0]['sourceText'],r'$$\frac{a_1}{b}=2$$')

    def test_local_layout_does_not_edit_embedded_formula_and_does_not_apply_to_inline(self):
        source=ole_fixture();p=profile_without_text_changes()
        result,_=transform(source,{'profile':p,'equations':{'o0':{'layout':{'align':'right','before':11}}}})
        a=parse(package(source)[0]['word/document.xml']);b=parse(package(result)[0]['word/document.xml'])
        self.assertEqual(E.tostring(a.find('.//'+q('object')),method='c14n'),E.tostring(b.find('.//'+q('object')),method='c14n'))
        self.assertEqual(val(child(child(b.find('.//'+q('p')),'pPr'),'jc')),'right')
        with self.assertRaises(ValueError):transform(source,{'profile':p,'equations':{'o0':{'format':{'size':12}}}})

    def test_project_restore_revalidates_formula_changes(self):
        import base64
        from server import restore_project, execute, import_document
        source=create_sample();p=profile_without_text_changes();item=import_document('自建公式示例.docx',source)
        settings={'id':item['id'],'profile':p,'equations':{'e0':{'format':{'size':14}}},'roles':{},'overrides':{},'tables':{},'insertions':[]}
        out,report=execute(settings)
        state={'format':'DocStyleProject','version':5,'name':'自建示例','profile':p,'documents':[{'id':item['id'],'name':item['name'],'sourceData':item['sourceData'],'sourceHash':item['sourceHash'],'equations':settings['equations'],'result':base64.b64encode(out).decode(),'report':report}],'workflow':{},'view':{}}
        restored=restore_project(state)['state']['documents'][0]
        self.assertEqual(restored['equations']['e0']['format']['size'],14)
        self.assertTrue(restored['report']['restoredAndRechecked'])
        self.assertIsNotNone(restored['result'])

class EquationIntegrationTests(unittest.TestCase):
    def test_preview_and_report_include_local_math_format_and_missing_font(self):
        from unittest.mock import patch
        from server import import_document, preview_changes, execute
        source=create_sample();p=profile_without_text_changes();d=import_document('自建示例.docx',source)
        settings={'id':d['id'],'profile':p,'roles':{},'overrides':{},'tables':{},'insertions':[],
                  'equations':{'e0':{'format':{'size':14,'font':'SelfBuilt Missing Math Font'}}}}
        fonts={'installedFonts':[],'commonFonts':[],'fonts':[],'source':'mock','detection':{'status':'mock'}}
        with patch('server.available_fonts',return_value=fonts):
            preview=preview_changes(settings);_,report=execute(settings)
        self.assertTrue(any(x['kind']=='equation' and x['field']=='size' and '覆盖全文' in x['source']['text'] for x in preview['changes']))
        self.assertIn('SelfBuilt Missing Math Font',report['fontAvailability']['notDetected'])
        self.assertGreater(report['fontAvailability']['mathRunsInspected'],0)
        self.assertEqual(next(x for x in report['checks'] if x['name']=='字体可用性')['status'],'pending')

    def test_requirements_parser_recognizes_mathematical_font(self):
        from rules import parse_local
        data=parse_local('公式字体 Cambria Math，字号12磅，居中')
        values={(x['target'],x['field']):x['value'] for x in data['items']}
        self.assertEqual(values[('equation','font')],'Cambria Math')
        self.assertEqual(values[('equation','size')],12)
        self.assertTrue(values[('equation','enabled')])
