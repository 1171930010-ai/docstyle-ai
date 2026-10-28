"""Functional regressions using real DOCX packages and python-docx readback."""
import copy
import io
import json
import unittest
import zipfile
from pathlib import Path
from docx import Document
from engine import default_profile, inspect_doc, transform, package, parse, q, val, child, formula, NS, style_context
from sample import create_sample

class FormattingTests(unittest.TestCase):
    def setUp(self):
        self.source=create_sample(); self.profile=default_profile(); self.model=inspect_doc(self.source)

    def effective(self,result,p):
        parts,_=package(result); return style_context(parts)[4](p._p,p.runs[0]._r if p.runs else None)

    def test_independent_heading_levels_and_paragraph_override(self):
        self.profile['styles']['heading1'].update(size=19,before=21,after=9,align='center',lineMode='exact',line=27,bold=False)
        self.profile['styles']['heading2'].update(size=13,before=7,after=4,align='right')
        body=next(p['id'] for p in self.model['paragraphs'] if p['text'].startswith('这段正文'))
        result,report=transform(self.source,{'profile':self.profile,'overrides':{body:{'before':18,'after':11,'left':.5,'indentMode':'hanging','indent':.4}}})
        d=Document(io.BytesIO(result)); h1=next(p for p in d.paragraphs if p.text=='一、研究设计');h2=next(p for p in d.paragraphs if p.text=='1.1 变量与样本')
        f1=self.effective(result,h1); f2=self.effective(result,h2)
        self.assertEqual(f1['size'],19);self.assertEqual(f1['before'],21)
        self.assertEqual(f1['after'],9);self.assertEqual(f1['line'],27);self.assertEqual(f1['lineMode'],'exact')
        self.assertFalse(f1['bold']);self.assertEqual(f2['size'],13)
        p=next(p for p in d.paragraphs if p.text.startswith('这段正文'))
        self.assertEqual(p.paragraph_format.space_before.pt,18);self.assertAlmostEqual(p.paragraph_format.first_line_indent.cm,-.4,2)
        self.assertTrue(next(r for r in p.runs if r.text=='原文文字不应改变').bold)
        self.assertEqual(report['checks'][-1]['status'],'pending')

    def test_complex_table_and_cell_override(self):
        self.profile['table'].update(enabled=True,headerRows=2,columns='3,3,3,4',rowHeight=22,paddingLeft=8)
        payload={'profile':self.profile,'tables':{'t0':{'cells':{'2:1':{'fill':'FFF1C7','bold':True,'align':'right','size':12,'borders':{'bottom':'double'}}}}}}
        result,report=transform(self.source,payload); d=Document(io.BytesIO(result)); t=d.tables[0]
        self.assertEqual(len(t.rows),5);self.assertEqual(t.cell(0,0)._tc,t.cell(1,0)._tc)
        self.assertEqual(t.cell(0,1)._tc,t.cell(0,2)._tc);self.assertEqual(t.cell(2,1).text,'22.14')
        self.assertTrue(t.cell(2,1).paragraphs[0].runs[0].bold);self.assertEqual(t.cell(2,1).paragraphs[0].runs[0].font.size.pt,12)
        self.assertEqual(val(child(child(t.cell(2,1)._tc,'tcPr'),'shd'),'fill'),'FFF1C7')
        self.assertTrue(t.rows[0]._tr.find('.//'+q('tblHeader')) is not None)
        self.assertEqual(report['counts']['tables'],1)

    def test_equation_edit_insert_and_untouched_formula(self):
        self.assertFalse(any(e['unsupported'] for e in self.model['equations']))
        anchor=next(p['id'] for p in self.model['paragraphs'] if p['text']=='二、模型设定')
        beforeparts,_=package(self.source); before=parse(beforeparts['word/document.xml'])
        result,report=transform(self.source,{'profile':self.profile,'equations':{'e0':{'latex':r'y_{it}=\alpha+\beta\frac{x_{it}}{z_{it}}'}},'insertions':[{'id':'n1','after':anchor,'latex':r'A=\begin{matrix}a&b\\c&d\end{matrix}'}]})
        afterparts,_=package(result); after=parse(afterparts['word/document.xml'])
        self.assertEqual(len(after.findall('.//'+q('m:oMath'))),3)
        self.assertEqual(before.xpath('//w:t/text()',namespaces=NS),after.xpath('//w:t/text()',namespaces=NS))
        self.assertEqual(report['counts']['equationsEdited'],1);self.assertEqual(report['counts']['equationsInserted'],1)
        self.assertIsNotNone(after.find('.//'+q('m:m')))
        self.assertTrue(all(v==afterparts[k] for k,v in beforeparts.items() if k not in report['allowedChangedParts']))
        Document(io.BytesIO(result))

    def test_disable_style_and_equation_formatting(self):
        for s in self.profile['styles'].values(): s['enabled']=False
        result,_=transform(self.source,{'profile':self.profile})
        parts,_=package(self.source);out,_=package(result)
        a=parse(parts['word/document.xml']); b=parse(out['word/document.xml'])
        from lxml import etree as E
        self.assertEqual(E.tostring(a,method='c14n'),E.tostring(b,method='c14n'))

    def test_all_nine_heading_styles(self):
        d=Document()
        for level in range(1,10): d.add_heading('标题 '+str(level),level=level)
        f=io.BytesIO();d.save(f)
        for level in range(1,10):self.profile['styles']['heading'+str(level)].update(size=20-level,before=level,after=level+1)
        result,_=transform(f.getvalue(),{'profile':self.profile})
        for level,p in enumerate(Document(io.BytesIO(result)).paragraphs,1):
            fmt=self.effective(result,p);self.assertEqual(fmt['size'],20-level);self.assertEqual(fmt['before'],level)
            self.assertEqual(val(child(child(p._p,'pPr'),'outlineLvl')),str(level-1))

    def test_page_and_equation_alignment(self):
        self.profile['page'].update(enabled=True,orientation='landscape',left=2,right=2,top=1.5)
        self.profile['equation'].update(enabled=True,align='right',before=8,after=10)
        result,_=transform(self.source,{'profile':self.profile});d=Document(io.BytesIO(result))
        self.assertGreater(d.sections[0].page_width,d.sections[0].page_height);self.assertAlmostEqual(d.sections[0].left_margin.cm,2,2)
        p=next(p for p in d.paragraphs if p._p.find(q('m:oMathPara')) is not None)
        self.assertEqual(p.paragraph_format.space_before.pt,8)
        self.assertEqual(p._p.find('.//'+q('m:jc')).get('{'+NS['m']+'}val'),'right')

    def test_invalid_inputs_fail_without_export(self):
        self.profile['styles']['body']['size']=-1
        with self.assertRaises(ValueError): transform(self.source,{'profile':self.profile})
        with self.assertRaises(ValueError):formula(r'\unknowncommand{x}')
        with self.assertRaises(ValueError):formula(r'\href{https://example.com}{x}')
        with self.assertRaises(Exception):inspect_doc(b'not a docx')

if __name__=='__main__':unittest.main()
