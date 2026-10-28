"""Self-built DOCX regressions for native styles, rule states and preservation."""
import copy
import base64
import io
import unittest
import zipfile
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from lxml import etree as E

import engine
from engine import NS, STYLE_IDS, child, default_profile, inspect_doc, package, parse, put, q, serialize, style_context, transform, val
from sample import create_sample


def saved(doc):
    out=io.BytesIO(); doc.save(out); return out.getvalue()


def replace_part(data,path,value):
    parts,infos=package(data); out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as z:
        for info in infos: z.writestr(info,value if info.filename==path else parts[info.filename])
    return out.getvalue()


class NativeStyleTests(unittest.TestCase):
    def test_bounded_format_retry_retains_direct_properties_and_all_guards(self):
        d=Document(); d.add_paragraph('标题'); p=d.add_heading('共同标题 ',1)
        p.add_run('强调文字').italic=True
        source=saved(d); profile=default_profile()
        profile['styles']['heading1'].update(size=17,font='黑体',before=8)
        result,report=transform(source,{'profile':profile,'_forceDirectFormatting':True})
        before,_=package(source); parts,_=package(result)
        original=parse(before['word/document.xml']); doc=parse(parts['word/document.xml'])
        self.assertEqual(engine.normalized_text(original),engine.normalized_text(doc))
        heading=doc.findall('.//'+q('p'))[1]
        self.assertEqual(val(child(child(heading,'pPr'),'pStyle')),STYLE_IDS['heading1'])
        self.assertEqual(val(child(child(heading,'pPr'),'spacing'),'before'),'160')
        for run in heading.findall(q('r')):
            self.assertEqual(val(child(child(run,'rPr'),'sz')),'34')
            self.assertEqual(val(child(child(run,'rPr'),'rFonts'),'eastAsia'),'黑体')
        effective=style_context(parts)[4]
        self.assertTrue(effective(heading,heading.findall(q('r'))[1])['italic'])
        self.assertGreaterEqual(report['validation']['textRunsChecked'],3)
        self.assertTrue(all(before[k]==parts[k] for k in before if k not in report['allowedChangedParts']))

    def test_word_global_style_edits_propagate_except_explicit_local_override(self):
        d=Document(); d.add_paragraph('标题'); h=d.add_heading('共同标题 ',1)
        inline=d.styles.add_style('SourceInlineTypography',WD_STYLE_TYPE.CHARACTER); inline.font.name='Courier New'; inline.font.size=Pt(31); inline.font.italic=True; inline.font.underline=True
        r=h.add_run('保留强调'); r.style=inline
        d.add_heading('另一共同标题',1); d.add_heading('局部覆盖标题',1)
        profile=default_profile(); profile['styles']['heading1'].update(size=17,bold=None)
        profile['styles']['heading1']['_fieldStates']={'size':'set','font':'set','latin':'set','italic':'preserve','underline':'preserve'}
        result,_=transform(saved(d),{'profile':profile,'overrides':{'p3':{'size':23,'font':'楷体'}}})
        parts,_=package(result); doc=parse(parts['word/document.xml']); sr=parse(parts['word/styles.xml']); role=next(s for s in sr.findall(q('style')) if val(s,'styleId')==STYLE_IDS['heading1'])
        # Simulate a real editor changing the common Word paragraph style.
        put(put(role,'rPr'),'sz',{'val':40}); put(child(role,'rPr'),'szCs',{'val':40})
        put(child(role,'rPr'),'rFonts',{'eastAsia':'黑体','ascii':'Arial','hAnsi':'Arial','cs':'Arial'})
        edited=replace_part(result,'word/styles.xml',serialize(sr)); newparts,_=package(edited); effective=style_context(newparts)[4]
        newdoc=parse(newparts['word/document.xml']); ps=newdoc.findall('.//'+q('p'))
        for p in ps[1:3]:
            for r in p.findall(q('r')):
                self.assertEqual(effective(p,r)['size'],20); self.assertEqual(effective(p,r)['font'],'黑体'); self.assertEqual(effective(p,r)['latin'],'Arial')
                self.assertIsNone(child(child(r,'rPr'),'sz'))
        self.assertTrue(effective(ps[1],ps[1].findall(q('r'))[1])['italic']); self.assertTrue(effective(ps[1],ps[1].findall(q('r'))[1])['underline'])
        self.assertEqual(effective(ps[3],ps[3].find(q('r')))['size'],23); self.assertEqual(effective(ps[3],ps[3].find(q('r')))['font'],'楷体')
        self.assertEqual(effective(ps[3],ps[3].find(q('r')))['latin'],'Arial')

    def test_preserved_heading_local_style_inherits_common_role_chain(self):
        d=Document(); d.add_paragraph('标题'); d.add_heading('第一节标题',1)
        profile=default_profile(); profile['styles']['heading1']['_fieldStates']={'font':'unspecified','size':'set','bold':'unspecified','color':'unspecified'}
        profile['styles']['heading1']['size']=17
        result,_=transform(saved(d),{'profile':profile}); parts,_=package(result); doc=parse(parts['word/document.xml']); sr=parse(parts['word/styles.xml'])
        nodes={val(s,'styleId'):s for s in sr.findall(q('style'))}; sid=val(child(child(doc.findall('.//'+q('p'))[1],'pPr'),'pStyle'))
        self.assertEqual(val(child(nodes[sid],'basedOn')),STYLE_IDS['heading1'])
        self.assertEqual(val(child(nodes[STYLE_IDS['heading1']],'basedOn')),STYLE_IDS['body'])
        self.assertEqual(val(child(nodes[STYLE_IDS['body']],'basedOn')),'Normal')
        self.assertEqual(inspect_doc(result)['paragraphs'][1]['format']['size'],17)

    def test_corrected_body_reimports_as_body_and_old_styles_survive(self):
        d=Document(); d.add_paragraph('研究标题',style='Title')
        d.add_paragraph('这是一段被错误设置为三级标题的完整正文。应该在再次导入时继续识别为正文。',style='Heading 3')
        source=saved(d); before,_=package(source)
        result,report=transform(source,{'profile':default_profile(),'roles':{'p1':'body'}})
        model=inspect_doc(result); self.assertEqual(model['paragraphs'][1]['role'],'body')
        out,_=package(result); root=parse(out['word/document.xml']); styles=parse(out['word/styles.xml'])
        body=root.findall('.//'+q('p'))[1]
        self.assertEqual(val(child(child(body,'pPr'),'pStyle')),STYLE_IDS['body'])
        self.assertEqual(val(child(child(body,'pPr'),'outlineLvl')),'9')
        original=parse(before['word/styles.xml'])
        for old in original.findall(q('style')):
            new=next(s for s in styles.findall(q('style')) if val(s,'styleId')==val(old,'styleId'))
            self.assertEqual(E.tostring(old,method='c14n'),E.tostring(new,method='c14n'))
        self.assertIn('word/styles.xml',report['allowedChangedParts'])
        self.assertTrue(all(before[k]==out[k] for k in before if k not in report['allowedChangedParts']))
        # Round-trip twice and keep the semantic correction.
        again,_=transform(result,{'profile':default_profile()})
        self.assertEqual(inspect_doc(again)['paragraphs'][1]['role'],'body')

    def test_native_role_styles_inheritance_continuation_and_outline(self):
        d=Document(); d.add_paragraph('论文标题'); d.add_paragraph('作者：张三'); d.add_paragraph('关键词：样式；继承')
        for level in range(1,10): d.add_heading('标题 '+str(level),level=level)
        result,_=transform(saved(d),{'profile':default_profile()})
        parts,_=package(result); sr=parse(parts['word/styles.xml']); nodes={val(n,'styleId'):n for n in sr.findall(q('style'))}
        for role,sid in STYLE_IDS.items():
            self.assertIn(sid,nodes); n=nodes[sid]
            self.assertEqual(val(n,'type'),'paragraph')
            self.assertEqual(val(child(n,'next')),STYLE_IDS['body'])
            self.assertEqual(val(child(n,'basedOn')),'Normal' if role=='body' else STYLE_IDS['body'])
            self.assertEqual(val(child(child(n,'pPr'),'outlineLvl')),str(int(role[7:])-1 if role.startswith('heading') else 9))
        model=inspect_doc(result)
        self.assertEqual([p['role'] for p in model['paragraphs'][:3]],['title','author','keywords'])
        self.assertEqual(model['paragraphs'][1]['format']['font'],'宋体')
        self.assertEqual(model['paragraphs'][1]['format']['align'],'center')

    def test_char_indent_is_independent_of_existing_centimeters(self):
        d=Document(); d.add_paragraph('标题'); p=d.add_paragraph('中文首行缩进正文。')
        put(put(p._p,'pPr'),'ind',{'firstLine':567,'hangingChars':300})
        profile=default_profile(); profile['styles']['body'].update(indent=2,indentUnit='chars')
        result,_=transform(saved(d),{'profile':profile})
        parts,_=package(result); p=parse(parts['word/document.xml']).findall('.//'+q('p'))[1]; ind=child(child(p,'pPr'),'ind')
        sr=parse(parts['word/styles.xml']); body=next(s for s in sr.findall(q('style')) if val(s,'styleId')==STYLE_IDS['body'])
        self.assertEqual(val(child(child(body,'pPr'),'ind'),'firstLineChars'),'200')
        for attr in ['firstLine','hanging','hangingChars']: self.assertIsNone(val(ind,attr))
        fmt=inspect_doc(result)['paragraphs'][1]['format']; self.assertEqual(fmt['indentUnit'],'chars'); self.assertEqual(fmt['indent'],2)
        profile['styles']['body'].update(indentMode='hanging',indent=1.5)
        result,_=transform(result,{'profile':profile}); parts,_=package(result)
        sr=parse(parts['word/styles.xml']); body=next(s for s in sr.findall(q('style')) if val(s,'styleId')==STYLE_IDS['body']); ind=child(child(body,'pPr'),'ind')
        self.assertEqual(val(ind,'hangingChars'),'150'); self.assertIsNone(val(ind,'firstLineChars'))

    def test_disabled_profile_has_no_native_style_changes(self):
        source=create_sample(); profile=default_profile()
        for s in profile['styles'].values(): s['enabled']=False
        result,report=transform(source,{'profile':profile}); before,_=package(source); after,_=package(result)
        self.assertEqual(before['word/styles.xml'],after['word/styles.xml'])
        self.assertNotIn('word/styles.xml',report['allowedChangedParts'])


class FieldStateTests(unittest.TestCase):
    def source(self):
        d=Document(); d.styles['Normal'].font.name='Arial'; d.styles['Normal'].font.size=Pt(11)
        style=d.styles.add_style('SourceBody',WD_STYLE_TYPE.PARAGRAPH); style.base_style=d.styles['Normal']
        style.font.name='Courier New'; style.font.size=Pt(17); style.paragraph_format.line_spacing=Pt(23)
        d.add_paragraph('标题',style='Title'); p=d.add_paragraph(style=style)
        p.add_run('普通文字'); r=p.add_run('强调文字'); r.bold=True; r.italic=True
        p.paragraph_format.space_before=Pt(8)
        put(put(p._p,'pPr'),'ind',{'firstLineChars':175})
        return saved(d)

    def test_preserve_and_unspecified_keep_each_run_effective_values(self):
        source=self.source(); profile=default_profile()
        profile['styles']['body'].update(font='黑体',size=26,bold=False,line=1.9,indent=0,
            _fieldStates={'font':'preserve','size':'preserve','bold':'preserve','italic':'unspecified','line':'preserve','indent':'preserve','before':'unspecified'})
        result,_=transform(source,{'profile':profile}); before,_=package(source); after,_=package(result)
        a=parse(before['word/document.xml']).findall('.//'+q('p'))[1]; b=parse(after['word/document.xml']).findall('.//'+q('p'))[1]
        old=style_context(before)[4]; new=style_context(after)[4]
        for x,y in zip(a.findall(q('r')),b.findall(q('r'))):
            for key in ['font','size','bold','italic','line','lineMode','indent','indentMode','indentUnit','before']:
                self.assertEqual(old(a,x)[key],new(b,y)[key],key)
        self.assertEqual(val(child(child(b,'pPr'),'ind'),'firstLineChars'),'175')

    def test_clear_removes_direct_format_and_resets_inheritance_deliberately(self):
        source=self.source(); profile=default_profile()
        profile['styles']['body'].update(font='黑体',size=26,bold=True,before=70,indent=2,indentUnit='chars',
            _fieldStates={'font':'clear','size':'clear','bold':'clear','before':'clear','indent':'clear'})
        result,_=transform(source,{'profile':profile}); out,_=package(result); p=parse(out['word/document.xml']).findall('.//'+q('p'))[1]
        pp=child(p,'pPr'); self.assertTrue(val(child(pp,'pStyle')).startswith(STYLE_IDS['body']+'__'))
        self.assertIsNone(val(child(pp,'spacing'),'before')); self.assertIsNone(val(child(pp,'ind'),'firstLineChars'))
        effective=style_context(out)[4]
        for r in p.findall(q('r')):
            rp=child(r,'rPr'); self.assertIsNone(child(rp,'sz')); self.assertIsNone(child(rp,'b'))
            self.assertIsNone(val(child(rp,'rFonts'),'eastAsia'))
            self.assertEqual(effective(p,r)['size'],11); self.assertFalse(effective(p,r)['bold'])
        self.assertEqual(effective(p)['before'],0); self.assertEqual(effective(p)['indent'],0)

    def test_local_set_overrides_preserve_without_polluting_global_template(self):
        profile=default_profile(); profile['styles']['body']['_fieldStates']={'size':'preserve'}
        result,_=transform(self.source(),{'profile':profile,'overrides':{'p1':{'size':21,'_fieldStates':{'size':'set'}}}})
        self.assertEqual(Document(io.BytesIO(result)).paragraphs[1].runs[0].font.size.pt,21)
        out,_=package(result); sr=parse(out['word/styles.xml']); body=next(n for n in sr.findall(q('style')) if val(n,'styleId')==STYLE_IDS['body'])
        self.assertIsNone(child(child(body,'rPr'),'sz'))

    def test_clear_preserves_unrelated_inherited_character_properties(self):
        d=Document(); d.add_paragraph('标题'); p=d.add_paragraph('正文 ')
        style=d.styles.add_style('InlineSource',WD_STYLE_TYPE.CHARACTER); style.font.bold=True; style.font.strike=True; style.font.superscript=True
        r=p.add_run('上标与删除线'); r.style=style
        profile=default_profile(); profile['styles']['body']['_fieldStates']={'bold':'clear'}; del profile['styles']['body']['bold']
        result,_=transform(saved(d),{'profile':profile}); out,_=package(result); doc=parse(out['word/document.xml']); r=doc.findall('.//'+q('p'))[1].findall(q('r'))[1]
        sid=val(child(child(r,'rPr'),'rStyle')); styles=parse(out['word/styles.xml']); sanitized=next(s for s in styles.findall(q('style')) if val(s,'styleId')==sid)
        rp=child(sanitized,'rPr'); self.assertIsNotNone(child(rp,'strike')); self.assertEqual(val(child(rp,'vertAlign')),'superscript'); self.assertIsNone(child(rp,'b'))
        p=doc.findall('.//'+q('p'))[1]; self.assertFalse(style_context(out)[4](p,r)['bold'])

    def test_manual_local_override_value_overrides_inherited_preserve_state(self):
        profile=default_profile(); profile['styles']['body']['_fieldStates']={'size':'preserve'}
        result,_=transform(self.source(),{'profile':profile,'overrides':{'p1':{'size':21}}})
        self.assertEqual(Document(io.BytesIO(result)).paragraphs[1].runs[0].font.size.pt,21)

    def test_reject_invalid_units_and_field_states(self):
        for changes in [{'_fieldStates':{'size':'bogus'}},{'indentUnit':'pixels'}]:
            profile=default_profile(); profile['styles']['body'].update(changes)
            with self.assertRaises(ValueError): transform(self.source(),{'profile':profile})


class FullReadbackTests(unittest.TestCase):
    def test_insertion_report_maps_original_paragraphs_to_new_positions(self):
        d=Document(); d.add_paragraph('标题'); d.add_paragraph('插入锚点。'); d.add_paragraph('后续正文保持正确定位。')
        result,report=transform(saved(d),{'profile':default_profile(),'insertions':[{'after':'p1','latex':'x=1'}]})
        self.assertEqual(report['paragraphMap'],{'p0':'p0','p1':'p1','p2':'p3'})
        model=inspect_doc(result); self.assertEqual(model['paragraphs'][3]['text'],'后续正文保持正确定位。')

    def test_square_root_has_required_degree_slot_and_hide_flag(self):
        root=engine.formula(r'\sqrt{x_1+x_2}')[0]; rad=root.find('.//'+q('m:rad'))
        self.assertIsNotNone(child(rad,'m:deg')); self.assertEqual(val(child(child(rad,'m:radPr'),'m:degHide'),'m:val'),'1')

    def test_picture_payload_drawing_relationships_and_text_survive(self):
        png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jChcAAAAASUVORK5CYII=')
        d=Document(); d.add_paragraph('标题'); p=d.add_paragraph('正文图片：'); p.add_run().add_picture(io.BytesIO(png)); p.add_run('图片后的文字。')
        source=saved(d); result,report=transform(source,{'profile':default_profile()}); before,_=package(source); after,_=package(result)
        self.assertEqual(inspect_doc(result)['counts']['images'],1)
        for path in before:
            if path.startswith('word/media/') or path.endswith('.rels'): self.assertEqual(before[path],after[path],path)
        a=parse(before['word/document.xml']); b=parse(after['word/document.xml'])
        self.assertEqual([E.tostring(n,method='c14n') for n in a.findall('.//'+q('drawing'))],[E.tostring(n,method='c14n') for n in b.findall('.//'+q('drawing'))])
        self.assertTrue(all(c['status']!='fail' for c in report['checks']))

    def test_all_runs_cell_runs_and_mixed_emphasis_are_checked(self):
        source=create_sample(); profile=default_profile(); profile['table']['enabled']=True
        result,report=transform(source,{'profile':profile})
        d=Document(io.BytesIO(result)); p=next(p for p in d.paragraphs if p.text.startswith('这段正文'))
        self.assertTrue(next(r for r in p.runs if r.text=='原文文字不应改变').bold)
        self.assertGreater(report['validation']['textRunsChecked'],report['counts']['paragraphs'])
        self.assertGreater(report['validation']['affectedParagraphs'],report['counts']['paragraphs'])

    def test_corruption_in_second_run_blocks_export(self):
        d=Document(); d.add_paragraph('标题'); p=d.add_paragraph(); p.add_run('第一段文字'); p.add_run('第二段文字')
        source=saved(d); real_serialize=engine.serialize
        def corrupt(root):
            clone=copy.deepcopy(root)
            if root.tag==q('document'):
                r=clone.findall('.//'+q('p'))[1].findall(q('r'))[1]
                put(put(r,'rPr'),'sz',{'val':99})
            return real_serialize(clone)
        with patch.object(engine,'serialize',side_effect=corrupt):
            with self.assertRaisesRegex(ValueError,'写回核验失败'): transform(source,{'profile':default_profile()})

    def test_postserialization_bookmark_corruption_blocks_export(self):
        d=Document(); p=d.add_paragraph('标题'); p._p.append(E.Element(q('bookmarkStart'),{q('id'):'8',q('name'):'bookmark'}))
        source=saved(d); real_serialize=engine.serialize
        def corrupt(root):
            clone=copy.deepcopy(root)
            if root.tag==q('document'): clone.find('.//'+q('bookmarkStart')).set(q('name'),'corrupted')
            return real_serialize(clone)
        with patch.object(engine,'serialize',side_effect=corrupt):
            with self.assertRaisesRegex(ValueError,'书签'): transform(source,{'profile':default_profile()})

    def test_bookmarks_fields_revisions_and_untouched_formula_survive(self):
        source=create_sample(); parts,_=package(source); root=parse(parts['word/document.xml']); p=root.findall('.//'+q('p'))[2]
        p.insert(1,E.Element(q('bookmarkStart'),{q('id'):'41',q('name'):'sample_anchor'})); p.append(E.Element(q('bookmarkEnd'),{q('id'):'41'}))
        r=E.SubElement(p,q('r')); E.SubElement(r,q('fldChar'),{q('fldCharType'):'begin'})
        r=E.SubElement(p,q('r')); E.SubElement(r,q('instrText')).text=' REF sample_anchor '
        r=E.SubElement(p,q('r')); E.SubElement(r,q('fldChar'),{q('fldCharType'):'end'})
        ins=E.SubElement(p,q('ins'),{q('id'):'9',q('author'):'Self-built test'}); r=E.SubElement(ins,q('r')); E.SubElement(r,q('t')).text='修订插入文字'
        deletion=E.SubElement(p,q('del'),{q('id'):'10',q('author'):'Self-built test'}); r=E.SubElement(deletion,q('r')); E.SubElement(r,q('delText')).text='保留删除文字'
        source=replace_part(source,'word/document.xml',serialize(root)); result,report=transform(source,{'profile':default_profile()}); after,_=package(result); final=parse(after['word/document.xml'])
        for tag in ['bookmarkStart','bookmarkEnd','fldChar','instrText','delText','m:oMath']:
            self.assertEqual([E.tostring(n,method='c14n') for n in root.findall('.//'+q(tag))],[E.tostring(n,method='c14n') for n in final.findall('.//'+q(tag))],tag)
        self.assertEqual(root.xpath('//w:t/text()',namespaces=NS),final.xpath('//w:t/text()',namespaces=NS))
        self.assertTrue(all(parts[k]==after[k] for k in parts if k not in report['allowedChangedParts']))


class LayoutStateTests(unittest.TestCase):
    def test_page_margin_preserve_and_clear(self):
        source=create_sample(); profile=default_profile(); profile['page'].update(enabled=True,top=1,left=1)
        profile['page']['_fieldStates']={'top':'preserve','left':'clear'}; del profile['page']['left']
        result,_=transform(source,{'profile':profile}); before,_=package(source); after,_=package(result)
        old=child(parse(before['word/document.xml']).find('.//'+q('sectPr')),'pgMar')
        new=child(parse(after['word/document.xml']).find('.//'+q('sectPr')),'pgMar')
        self.assertEqual(val(old,'top'),val(new,'top')); self.assertIsNone(val(new,'left'))

    def test_equation_layout_preserve_and_clear_without_missing_keys(self):
        source=create_sample(); profile=default_profile(); profile['equation'].update(enabled=True,align='right')
        profile['equation']['_fieldStates']={'align':'preserve','before':'clear'}; del profile['equation']['before']
        result,_=transform(source,{'profile':profile}); before,_=package(source); after,_=package(result)
        old=parse(before['word/document.xml']); new=parse(after['word/document.xml'])
        self.assertEqual([E.tostring(n,method='c14n') for n in old.findall('.//'+q('m:jc'))],[E.tostring(n,method='c14n') for n in new.findall('.//'+q('m:jc'))])
        for p in new.findall('.//'+q('p')):
            if child(p,'m:oMathPara') is not None: self.assertIsNone(val(child(child(p,'pPr'),'spacing'),'before'))
        profile['equation']['_fieldStates']['align']='clear'; del profile['equation']['align']
        result,_=transform(result,{'profile':profile}); after,_=package(result)
        self.assertFalse(parse(after['word/document.xml']).findall('.//'+q('m:oMathParaPr')+'/'+q('m:jc')))

    def test_table_typography_clear_and_layout_states_fail_explicitly(self):
        source=create_sample(); profile=default_profile(); profile['table']['enabled']=True
        profile['table']['_fieldStates']={'size':'clear','font':'preserve'}; del profile['table']['size']
        result,_=transform(source,{'profile':profile}); out,_=package(result); root=parse(out['word/document.xml'])
        for r in root.xpath('//w:tc/w:p/w:r',namespaces=NS): self.assertIsNone(child(child(r,'rPr'),'sz'))
        profile['table']['_fieldStates']['width']='clear'
        with self.assertRaisesRegex(ValueError,'表格布局字段'): transform(source,{'profile':profile})

    def test_partial_table_size_rule_preserves_all_unconfirmed_layout(self):
        d=Document(); d.add_paragraph('标题'); t=d.add_table(rows=2,cols=2); t.style='Table Grid'
        for cell in t.rows[0].cells+t.rows[1].cells:
            p=cell.paragraphs[0]; p.add_run('自建单元格').font.size=Pt(12)
            put(put(p._p,'pPr'),'jc',{'val':'right'}); put(put(p._p,'pPr'),'ind',{'firstLineChars':100})
        put(put(t.rows[0]._tr,'trPr'),'trHeight',{'val':460,'hRule':'exact'})
        source=saved(d); profile=default_profile()
        for s in profile['styles'].values(): s['enabled']=False
        profile['table']['enabled']=True; profile['table']['size']=14
        profile['table']['_fieldStates']={k:'unspecified' for k in profile['table'] if k not in ['cells','_fieldStates']}
        profile['table']['_fieldStates'].update(enabled='set',size='set')
        result,_=transform(source,{'profile':profile}); before,_=package(source); after,_=package(result)
        a=parse(before['word/document.xml']).find('.//'+q('tbl')); b=parse(after['word/document.xml']).find('.//'+q('tbl'))
        for tag in ['tblPr','tblGrid','tcPr','trPr','pPr']:
            self.assertEqual([E.tostring(n,method='c14n') for n in a.findall('.//'+q(tag))],[E.tostring(n,method='c14n') for n in b.findall('.//'+q(tag))],tag)
        for r in b.findall('.//'+q('r')): self.assertEqual(val(child(child(r,'rPr'),'sz')),'28')

    def test_enabled_fully_unspecified_table_changes_nothing(self):
        source=create_sample(); profile=default_profile()
        for s in profile['styles'].values(): s['enabled']=False
        profile['table']['enabled']=True
        profile['table']['_fieldStates']={k:'unspecified' for k in profile['table'] if k!='_fieldStates'}
        result,report=transform(source,{'profile':profile}); before,_=package(source); after,_=package(result)
        self.assertEqual(before,after); self.assertEqual(report['counts']['tables'],0)

    def test_partial_table_fixed22pt_rule_preserves_all_other_layout(self):
        d=Document(); d.add_paragraph('标题'); table=d.add_table(rows=1,cols=1); table.cell(0,0).text='单元格正文'
        p=table.cell(0,0).paragraphs[0]; put(put(p._p,'pPr'),'jc',{'val':'right'}); put(put(p._p,'pPr'),'ind',{'firstLineChars':100})
        source=saved(d); profile=default_profile()
        for s in profile['styles'].values(): s['enabled']=False
        profile['table'].update(enabled=True,line=22,lineMode='exact')
        profile['table']['_fieldStates']={k:'unspecified' for k in profile['table'] if k!='_fieldStates'}
        profile['table']['_fieldStates'].update(enabled='set',line='set',lineMode='set')
        result,_=transform(source,{'profile':profile}); before,_=package(source); after,_=package(result)
        a=parse(before['word/document.xml']).find('.//'+q('tbl')); b=parse(after['word/document.xml']).find('.//'+q('tbl'))
        for tag in ['tblPr','tblGrid','tcPr','trPr']:
            self.assertEqual([E.tostring(n,method='c14n') for n in a.findall('.//'+q(tag))],[E.tostring(n,method='c14n') for n in b.findall('.//'+q(tag))],tag)
        pp=child(b.find('.//'+q('p')),'pPr'); self.assertEqual(val(child(pp,'jc')),'right'); self.assertEqual(val(child(pp,'ind'),'firstLineChars'),'100')
        self.assertEqual(val(child(pp,'spacing'),'line'),'440'); self.assertEqual(val(child(pp,'spacing'),'lineRule'),'exact')

    def test_fully_unspecified_page_and_equation_do_not_add_format_nodes(self):
        source=create_sample(); profile=default_profile()
        for s in profile['styles'].values(): s['enabled']=False
        for target in ['page','equation']:
            profile[target]['enabled']=True; profile[target]['_fieldStates']={k:'unspecified' for k in profile[target] if k!='_fieldStates'}
        result,report=transform(source,{'profile':profile}); before,_=package(source); after,_=package(result)
        self.assertEqual(before,after); self.assertEqual(report['counts']['sections'],0); self.assertEqual(report['counts']['equationParagraphs'],0)

    def test_page_unspecified_geometry_preserves_custom_dimensions(self):
        d=Document(); d.add_paragraph('标题'); section=d.sections[0]; section.page_width=Pt(500); section.page_height=Pt(700)
        source=saved(d); profile=default_profile(); profile['page'].update(enabled=True,left=1.8)
        profile['page']['_fieldStates']={k:'unspecified' for k in profile['page']}; profile['page']['_fieldStates'].update(enabled='set',left='set')
        result,_=transform(source,{'profile':profile}); a=Document(io.BytesIO(source)); b=Document(io.BytesIO(result))
        self.assertEqual(a.sections[0].page_width,b.sections[0].page_width); self.assertEqual(a.sections[0].page_height,b.sections[0].page_height)


if __name__=='__main__': unittest.main()
