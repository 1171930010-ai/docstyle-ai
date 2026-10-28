# coding: utf-8
"""Real loopback HTTP plus injectable mock-provider regression evidence."""
import unittest,threading,urllib.request,urllib.error,json,base64,tempfile,copy,time,io,zipfile
from unittest.mock import patch
from docx import Document
from docx.shared import Pt
import server
from project_store import LocalStore
from ai_pipeline import analyze_document,parse_requirements
from engine import inspect_doc,package

class ProviderTests(unittest.TestCase):
 def document(self):
  d=Document();d.add_paragraph('测试标题')
  for i in range(92):d.add_paragraph('这是正文完整句子，用来检查跨块结构识别不能丢失段落。',style='Heading 3' if i==40 else None)
  f=io.BytesIO();d.save(f);return inspect_doc(f.getvalue())
 def test_shared_full_outline_chunks_and_locks(self):
  model=self.document();calls=[]
  def request(key,model,messages,structured,tokens):
   data=json.loads(messages[-1]['content']);calls.append(data)
   if 'document_skeleton' in data:return json.dumps({'outline':[{'id':'p0','role':'title'}],'sections':[]})
   if 'review_stage' in data:return json.dumps({'corrections':[],'warnings':['模拟跨块复查']})
   return json.dumps({'items':[{'id':p['id'],'role':'body','confidence':.7,'reason':'模拟结果，仅测流程'} for p in data['paragraphs']]})
  out=analyze_document(request,'test-only-not-a-key','mock',model,server.confirmed_defaults(),locks={'p0':True},roles={'p0':'title'})
  self.assertEqual(out['total'],93);self.assertEqual(out['omitted'],0);self.assertEqual(out['chunks'],3);self.assertTrue(out['globalReviewPerformed']);self.assertEqual(out['items'][0]['role'],'title')
  for c in calls[1:]:self.assertEqual(c['global_context']['ai_outline'][0]['id'],'p0')
  self.assertEqual(out['items'][41]['role'],'body')
 def test_missing_ids_are_retried_once_then_left_pending(self):
  calls=[]
  def request(key,model,messages,structured,tokens):
   calls.append(messages);return '{"outline":[],"items":[]}'
  out=analyze_document(request,'mock','mock',self.document(),server.confirmed_defaults())
  self.assertEqual(out['retryCount'],1);self.assertEqual(out['omitted'],93);self.assertGreater(len(out['issues']),0);self.assertEqual(len(calls),7)
 def test_rules_provider_does_not_auto_confirm(self):
  raw={'items':[{'target':'style','role':'body','field':'size','state':'set','value':12,'unit':'pt','status':'confirmed'}],'warnings':['模拟API接口结果']}
  out=parse_requirements(lambda *a:json.dumps(raw),'mock','mock','正文小四')
  self.assertEqual(out['items'][0]['status'],'pending');self.assertEqual(out['items'][0]['source']['originalRequirements'],'正文小四')

class HTTPTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.temp=tempfile.TemporaryDirectory();cls.store=server.DISK;server.DISK=LocalStore(cls.temp.name)
  cls.fontpatch=patch.object(server,'available_fonts',return_value={'installedFonts':['Arial'],'commonFonts':['宋体'],'fonts':['Arial'],'source':'模拟安装目录','detection':{'status':'mock','renderVerified':False}});cls.fontpatch.start()
  cls.http=server.ThreadingHTTPServer(('127.0.0.1',0),server.Handler);cls.port=cls.http.server_port;threading.Thread(target=cls.http.serve_forever,daemon=True).start()
 @classmethod
 def tearDownClass(cls):cls.http.shutdown();cls.http.server_close();cls.fontpatch.stop();server.DISK=cls.store;cls.temp.cleanup()
 def request(self,path,payload=None,origin=None):
  headers={'Content-Type':'application/json'}
  if origin:headers['Origin']=origin
  req=urllib.request.Request(f'http://127.0.0.1:{self.port}/api/'+path,data=json.dumps(payload).encode() if payload is not None else None,headers=headers)
  with urllib.request.urlopen(req,timeout=20) as r:return json.load(r)
 def sample(self):return self.request('sample')
 def settings(self,d):return {'id':d['id'],'profile':self.request('defaults'),'roles':{},'overrides':{},'tables':{},'equations':{},'insertions':[]}
 def test_http_assets_and_origin(self):
  self.assertEqual(self.request('health')['version'],'5.0')
  with self.assertRaises(urllib.error.HTTPError) as c:self.request('health',origin='https://external.invalid')
  self.assertEqual(c.exception.code,403)
  with urllib.request.urlopen(f'http://127.0.0.1:{self.port}/workflow.js') as r:self.assertEqual(r.status,200)
 def test_partial_rules_preserve_unspecified_body_fields(self):
  d=Document();p=d.add_paragraph('完整正文句子。');p.paragraph_format.space_after=Pt(17);p.runs[0].font.size=Pt(20);f=io.BytesIO();d.save(f)
  imp=self.request('import',{'name':'参数.docx','data':base64.b64encode(f.getvalue()).decode()});settings=self.settings(imp)
  parsed=self.request('rules/parse',{'text':'正文段前0磅','mode':'local'})
  for x in parsed['items']:x['status']='confirmed'
  settings['profile']=self.request('rules/apply',{'profile':settings['profile'],'items':parsed['items']})['profile']
  out=self.request('format',settings);p=inspect_doc(base64.b64decode(out['data']))['paragraphs'][0];self.assertEqual(p['format']['after'],17);self.assertEqual(p['format']['size'],20);self.assertEqual(p['format']['before'],0)
 def test_missing_font_is_reported_from_every_output_run(self):
  d=Document();p=d.add_paragraph('字体检查。');r=p.add_run('第二run');r.font.name='Nonexistent_Font_SelfBuilt_Test';f=io.BytesIO();d.save(f)
  imp=self.request('import',{'name':'自建缺失字体.docx','data':base64.b64encode(f.getvalue()).decode()});out=self.request('format',self.settings(imp))
  self.assertIn('Nonexistent_Font_SelfBuilt_Test',out['report']['fontAvailability']['notDetected']);self.assertEqual(next(x for x in out['report']['checks'] if x['name']=='字体可用性')['status'],'pending')
 def test_manual_profile_numeric_boolean_is_rejected(self):
  d=self.sample();s=self.settings(d);s['profile']['styles']['body']['size']=True;s['profile']['styles']['body']['_fieldStates']['size']='set'
  with self.assertRaises(urllib.error.HTTPError):self.request('format',s)
 def test_unknown_ai_key_is_pending(self):
  with self.assertRaises(urllib.error.HTTPError) as c:self.request('rules/parse',{'text':'正文宋体','mode':'deepseek'})
  self.assertIn('尚未验证',json.load(c.exception)['error'])
 def test_preview_and_corrected_reimport(self):
  d=self.sample();s=self.settings(d);local=self.request('structure',{**s,'locks':{},'roles':{}});s['roles']={x['id']:x['role'] for x in local['items']}
  preview=self.request('preview',s);self.assertGreater(len(preview['changes']),0);out=self.request('format',s);self.assertTrue(all(c['status']!='fail' for c in out['report']['checks']))
  re=self.request('import',{'name':'重新导入.docx','data':out['data']});self.assertEqual([p['text'] for p in re['model']['paragraphs']],[p['text'] for p in d['model']['paragraphs']])
 def test_style_repair_is_bounded_and_protected_failures_are_not_retried(self):
  d=self.sample();s=self.settings(d);actual=server.transform;calls=[]
  def once(data,payload):
   calls.append(payload)
   if len(calls)==1:raise ValueError('段落格式写回核验失败：模拟首轮继承异常')
   return actual(data,payload)
  with patch.object(server,'transform',side_effect=once):out=self.request('format',s)
  self.assertEqual(len(calls),2);self.assertTrue(calls[1]['_forceDirectFormatting']);self.assertEqual(len(out['report']['retryHistory']),1)
  with patch.object(server,'transform',side_effect=ValueError('原文文字一致性校验失败')) as mocked:
   with self.assertRaises(urllib.error.HTTPError):self.request('format',s)
   self.assertEqual(mocked.call_count,1)
  with patch.object(server,'transform',side_effect=ValueError('段落格式写回核验失败：模拟永久失败')) as mocked:
   with self.assertRaises(urllib.error.HTTPError) as error:self.request('format',s)
   self.assertEqual(mocked.call_count,2);self.assertIn('仍失败',json.load(error.exception)['error'])
 def test_preview_insertion_does_not_shift_original_paragraphs(self):
  d=self.sample();s=self.settings(d);s['insertions']=[{'id':'new','after':'p0','latex':'x+y'}]
  out=self.request('preview',s)
  # Unknown font/size is preserved for corresponding paragraphs even after insertion.
  self.assertFalse(any(x['field']=='size' for x in out['changes']));self.assertTrue(any(x['field']=='insert' for x in out['changes']))
 def test_project_restore_rechecks_output_and_recreates_ids(self):
  d=self.sample();s=self.settings(d);s['roles']={'p0':'title'};out=self.request('format',s)
  record={**d,**s,'sourceData':d['sourceData'],'sourceHash':d['sourceHash'],'aiLocks':{'p0':True},'result':out['data'],'report':{'checks':[{'name':'FAKE','status':'pass'}],'pagination':{'sourceSHA256':out['report']['resultSHA256'],'id':'old-session','pageCount':1,'WordWpsVerified':True}}}
  state={'profile':s['profile'],'documents':[record],'workflow':{'pdfs':[],'apiKey':'not-a-real-key'},'view':{'current':d['id']}}
  saved=self.request('projects/save',{'name':'恢复','state':state});restored=self.request('projects/restore',{'id':saved['project']['id']});r=restored['state']['documents'][0]
  self.assertNotEqual(r['id'],d['id']);self.assertEqual(restored['state']['view']['current'],r['id']);self.assertTrue(r['report']['restoredAndRechecked']);self.assertFalse(r['report']['pagination']['WordWpsVerified']);self.assertTrue(r['report']['pagination']['restoredMetadataUnverified']);self.assertNotIn('id',r['report']['pagination']);self.assertNotIn('apiKey',json.dumps(restored));self.assertEqual(r['aiLocks'],{'p0':True})
 def test_restored_untrusted_output_is_not_marked_pass(self):
  d=self.sample();s=self.settings(d);state={'profile':s['profile'],'documents':[{**d,'result':base64.b64encode(b'invalid zip').decode(),'report':{'checks':[{'status':'pass'}]}}],'workflow':{'pdfs':[]}}
  out=self.request('projects/restore',{'state':state})['state']['documents'][0];self.assertIsNone(out['result']);self.assertIsNone(out['report']);self.assertIn('未通过',out['restoreWarning'])
 def test_batch_partial_failure_then_repair_retry_packages_all(self):
  d=self.sample();s=self.settings(d);failed={**s,'id':'test-late-import'}
  job=self.request('tasks/start',{'documents':[s,failed]})['id']
  def wait(id):
   for _ in range(100):
    x=self.request('tasks/'+id)
    if x['status'] not in ['queued','running','cancelling']:return x
    time.sleep(.02)
   self.fail('batch timeout')
  first=wait(job);self.assertEqual(first['status'],'completed_with_errors')
  # Simulate recovery of lost source session; no mocked DOCX execution.
  with server.LOCK:server.STORE['test-late-import']=server.STORE[d['id']].copy()
  retry=self.request('tasks/retry',{'id':job})['id'];second=wait(retry);self.assertEqual(second['progress']['success'],2);self.assertEqual(second['items'][0]['attempts'],1)
  with urllib.request.urlopen(f'http://127.0.0.1:{self.port}/api/tasks/{retry}/download') as r:z=zipfile.ZipFile(io.BytesIO(r.read()))
  self.assertEqual(len(z.namelist()),5)
 def test_pdf_page_evidence_and_template_version(self):
  from pathlib import Path
  file=next(Path('validation/v5/pdf_fixtures').glob('*digital*.pdf'),None) or next(Path('validation/v5/pdf_fixtures').glob('*.pdf'))
  p=self.request('pdf/import',{'name':file.name,'data':base64.b64encode(file.read_bytes()).decode(),'metadata':{'journal':'自建测试'}})
  self.assertIn(p['status'],['pending_confirmation','needs_ocr']);self.assertEqual(p['sourceHash'],__import__('hashlib').sha256(file.read_bytes()).hexdigest())
  with urllib.request.urlopen(f'http://127.0.0.1:{self.port}/api/pdf/{p["id"]}/page/1') as r:self.assertEqual(r.read(8),b'\x89PNG\r\n\x1a\n')
  candidates=[x for x in p['candidates'] if x['state']=='set']
  if candidates:
   item=candidates[0];item['status']='confirmed';a=self.request('templates/save',{'name':'PDF测试','purpose':'publication_reference','metadata':p['metadata'],'items':[item],'pdfIds':[p['id']]})['template'];b=self.request('templates/save',{'name':'PDF测试','purpose':'publication_reference','metadata':p['metadata'],'items':[item],'pdfIds':[p['id']]})['template'];self.assertEqual(b['version'],a['version']+1)
if __name__=='__main__':unittest.main()
